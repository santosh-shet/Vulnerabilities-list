#!/usr/bin/env python3
"""
Build the daily Excel and PDF reports plus the GitHub Pages dashboard
from the most recent snapshot in data/.

Usage:
    python scripts/build_reports.py                 # latest snapshot
    python scripts/build_reports.py --date 2026-09-10
"""
import argparse
import html
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, TableStyle
from reportlab.platypus import Table as RLTable

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
REPORT_DIR = ROOT / "reports"
DOCS_DIR = ROOT / "docs"

SEVERITY_FILL = {
    "CRITICAL": "C00000",
    "HIGH": "E97132",
    "MEDIUM": "F2C94C",
    "LOW": "8FBC8F",
}
SEVERITY_FONT = {"CRITICAL": "FFFFFF", "HIGH": "FFFFFF", "MEDIUM": "000000", "LOW": "000000"}

COLUMNS = [
    ("Priority", "priority", 16),
    ("Reason", "reason", 13),
    ("ID", "cve_id", 18),
    ("CVSS", "cvss_score", 7),
    ("Severity", "cvss_severity", 10),
    ("Score src", "cvss_source", 10),
    ("KEV", "kev", 6),
    ("EPSS", "epss", 7),
    ("SSVC", "ssvc_exploitation", 9),
    ("Watchlist", "watchlist_hits", 16),
    ("Vendor / Product", "vendor_product", 28),
    ("Package", "ecosystem_package", 28),
    ("Fixed in", "patched_version", 14),
    ("Sources", "sources", 18),
    ("Published", "published", 12),
    ("KEV Due", "kev_due_date", 11),
    ("CWE", "cwe", 12),
    ("Description", "description", 90),
    ("Link", "nvd_url", 42),
]


def load_snapshot(date: str | None):
    files = sorted(DATA_DIR.glob("*.json"))
    if not files:
        raise SystemExit("No snapshots in data/. Run fetch_vulns.py first.")
    path = DATA_DIR / f"{date}.json" if date else files[-1]
    with open(path) as f:
        return json.load(f)


# Characters openpyxl/Excel refuse in a cell: control chars except tab/newline/CR
_ILLEGAL_XLSX = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def clean(s: str) -> str:
    return _ILLEGAL_XLSX.sub("", s)


def fmt(v):
    if v is None:
        return ""
    if isinstance(v, list):
        return clean(", ".join(str(x) for x in v))
    if isinstance(v, bool):
        return "Yes" if v else ""
    if isinstance(v, float):
        return round(v, 3)
    if isinstance(v, str) and len(v) > 10 and v[4] == "-" and "T" in v:
        return v[:10]  # ISO datetime -> date
    if isinstance(v, str):
        return clean(v)
    return v


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def write_sheet(ws, records, title):
    ws.title = title[:31]
    headers = [c[0] for c in COLUMNS]
    ws.append(headers)
    for r in records:
        ws.append([fmt(r.get(c[1])) for c in COLUMNS])

    hdr_font = Font(name="Arial", bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    for cell in ws[1]:
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for i, (_, _, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    # Hyperlinks on the NVD column
    link_col = [c[0] for c in COLUMNS].index("Link") + 1
    for idx, r in enumerate(records, start=2):
        cell = ws.cell(row=idx, column=link_col)
        if not cell.value:
            cell.value = r.get("advisory_url") or ""
        if cell.value:
            cell.hyperlink = cell.value
            cell.font = Font(name="Arial", size=10, color="0563C1", underline="single")

    # Conditional formatting on Severity
    if ws.max_row > 1:
        sev_col = get_column_letter([c[0] for c in COLUMNS].index("Severity") + 1)
        rng = f"{sev_col}2:{sev_col}{ws.max_row}"
        for sev, fill in SEVERITY_FILL.items():
            ws.conditional_formatting.add(
                rng,
                CellIsRule(
                    operator="equal",
                    formula=[f'"{sev}"'],
                    fill=PatternFill("solid", fgColor=fill),
                    font=Font(color=SEVERITY_FONT[sev], bold=True),
                ),
            )
        kev_col = get_column_letter([c[0] for c in COLUMNS].index("KEV") + 1)
        ws.conditional_formatting.add(
            f"{kev_col}2:{kev_col}{ws.max_row}",
            CellIsRule(
                operator="equal",
                formula=['"Yes"'],
                fill=PatternFill("solid", fgColor="7030A0"),
                font=Font(color="FFFFFF", bold=True),
            ),
        )


def build_xlsx(snap, out: Path):
    recs = snap["records"]
    wb = Workbook()

    # Summary sheet
    ws = wb.active
    ws.title = "Summary"
    c = snap["counts"]
    rows = [
        ("Daily Vulnerability Report", ""),
        ("Report date", snap["report_date"]),
        ("Window (UTC)", f"{snap['window_start'][:16]} to {snap['window_end'][:16]}"),
        ("Generated", snap["generated_at"][:19].replace("T", " ") + " UTC"),
        ("", ""),
        ("Total entries", c["total"]),
        ("  Newly published", c.get("new", "")),
        ("  Updated (recently published, now scored / KEV)", c.get("updated", "")),
        ("In CISA KEV (actively exploited)", c["kev"]),
        ("Critical", c["critical"]),
        ("High", c["high"]),
        ("Medium", c["medium"]),
        ("Low", c["low"]),
        ("Unscored (awaiting NVD analysis)", c["unscored"]),
        ("Watchlist matches", c["watchlist"]),
        ("From GitHub Advisories", c.get("ghsa", "")),
        ("From OSV.dev", c.get("osv", "")),
        ("", ""),
        ("Sources", "NVD 2.0, CISA KEV, FIRST EPSS, CISA Vulnrichment, GitHub Security Advisories, OSV.dev"),
        ("Watchlist keywords", ", ".join(snap.get("watchlist_keywords", [])) or "(none configured)"),
        ("", ""),
        ("Priority rules", ""),
        ("P1 - Act now", "In CISA KEV, or EPSS >= 0.50, or SSVC exploitation = active"),
        ("P2 - High", "CVSS >= 9.0, or CVSS >= 7.0 with EPSS >= 0.10, or SSVC exploitation = poc"),
        ("P3 - Medium", "CVSS >= 7.0"),
        ("P4 - Low / Unscored", "Everything else, incl. CVEs not yet scored by NVD"),
    ]
    for r in rows:
        ws.append(list(r))
    ws["A1"].font = Font(name="Arial", size=14, bold=True)
    for row in ws.iter_rows(min_row=2):
        row[0].font = Font(name="Arial", bold=True)
        row[1].font = Font(name="Arial")
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 70

    # Detail sheets
    write_sheet(wb.create_sheet(), recs, "All CVEs")
    write_sheet(wb.create_sheet(), [r for r in recs if r["kev"]], "KEV - Exploited")
    write_sheet(
        wb.create_sheet(),
        [r for r in recs if r["cvss_severity"] in ("CRITICAL", "HIGH")],
        "Critical & High",
    )
    if snap.get("watchlist_keywords"):
        write_sheet(wb.create_sheet(), [r for r in recs if r["watchlist_match"]], "Watchlist")
    write_sheet(wb.create_sheet(), [r for r in recs if r.get("ecosystem_package")], "Packages (GHSA-OSV)")

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


# --------------------------------------------------------------------------- #
# PDF (executive summary + P1/P2 detail)
# --------------------------------------------------------------------------- #
def build_pdf(snap, out: Path):
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["Normal"], fontName="Helvetica", fontSize=8, leading=10)
    small = ParagraphStyle("small", parent=body, fontSize=7, leading=9, textColor=colors.grey)
    h1 = styles["Title"]
    h2 = styles["Heading2"]

    doc = SimpleDocTemplate(
        str(out),
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=f"Daily Vulnerability Report {snap['report_date']}",
    )
    c = snap["counts"]
    story = [
        Paragraph(f"Daily Vulnerability Report - {snap['report_date']}", h1),
        Paragraph(
            f"Window (UTC): {snap['window_start'][:16]} to {snap['window_end'][:16]} &nbsp;|&nbsp; "
            f"Sources: NVD, CISA KEV, EPSS, Vulnrichment, GHSA, OSV &nbsp;|&nbsp; "
            f"Generated {snap['generated_at'][:19].replace('T', ' ')} UTC",
            small,
        ),
        Spacer(1, 6),
    ]

    summary = [
        ["Total", "New", "Updated", "KEV", "Critical", "High", "Medium", "Unscored", "Watchlist", "Packages"],
        [c["total"], c.get("new", ""), c.get("updated", ""), c["kev"], c["critical"], c["high"], c["medium"], c["unscored"], c["watchlist"], c.get("ghsa", 0)],
    ]
    t = RLTable(summary, hAlign="LEFT")
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("BACKGROUND", (3, 1), (3, 1), colors.HexColor("#E6D9F2")),
                ("BACKGROUND", (4, 1), (4, 1), colors.HexColor("#F4CCCC")),
                ("BACKGROUND", (5, 1), (5, 1), colors.HexColor("#FCE5CD")),
            ]
        )
    )
    story += [t, Spacer(1, 10)]

    def detail_table(records, heading):
        story.append(Paragraph(heading, h2))
        if not records:
            story.append(Paragraph("None in this window.", body))
            story.append(Spacer(1, 6))
            return
        data = [["ID", "CVSS", "Sev", "KEV", "EPSS", "Vendor / Product / Package", "Description"]]
        for r in records:
            desc = r["description"]
            if len(desc) > 420:
                desc = desc[:417] + "..."
            data.append(
                [
                    Paragraph(f'<a href="{r.get("nvd_url") or r.get("advisory_url") or "#"}" color="blue">{r["cve_id"]}</a>'
                              f'<br/><font size="6" color="grey">{html.escape(r.get("reason",""))}</font>', body),
                    fmt(r["cvss_score"]),
                    r["cvss_severity"][:4],
                    "Yes" if r["kev"] else "",
                    fmt(r["epss"]),
                    Paragraph(html.escape((r["vendor_product"] or r.get("ecosystem_package", ""))[:120])
                              + (f'<br/><font size="6" color="grey">fixed: {html.escape(r["patched_version"][:60])}</font>' if r.get("patched_version") else ""), body),
                    Paragraph(html.escape(desc), body),
                ]
            )
        tbl = RLTable(
            data,
            colWidths=[28 * mm, 12 * mm, 12 * mm, 10 * mm, 12 * mm, 50 * mm, 143 * mm],
            repeatRows=1,
        )
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F7F7")]),
        ]
        for i, r in enumerate(records, start=1):
            sev = r["cvss_severity"]
            if sev in SEVERITY_FILL:
                style.append(("BACKGROUND", (2, i), (2, i), colors.HexColor("#" + SEVERITY_FILL[sev])))
                style.append(("TEXTCOLOR", (2, i), (2, i), colors.HexColor("#" + SEVERITY_FONT[sev])))
            if r["kev"]:
                style.append(("BACKGROUND", (3, i), (3, i), colors.HexColor("#7030A0")))
                style.append(("TEXTCOLOR", (3, i), (3, i), colors.white))
        tbl.setStyle(TableStyle(style))
        story.append(tbl)
        story.append(Spacer(1, 8))

    recs = snap["records"]
    detail_table([r for r in recs if r["priority"].startswith("P1")], "P1 - Act now (KEV / high EPSS)")
    detail_table([r for r in recs if r["priority"].startswith("P2")], "P2 - High (CVSS >= 9.0 or 7.0+ with EPSS >= 0.10)")
    if snap.get("watchlist_keywords"):
        detail_table(
            [r for r in recs if r["watchlist_match"] and not r["priority"].startswith(("P1", "P2"))],
            "Watchlist matches (other priorities)",
        )
    story.append(
        Paragraph(
            f"P3/P4 items ({sum(r['priority'].startswith(('P3', 'P4')) for r in recs)}) are in the Excel report. "
            "Unscored CVEs are newly published and awaiting NVD analysis; re-check tomorrow's report.",
            small,
        )
    )
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.build(story)


# --------------------------------------------------------------------------- #
# GitHub Pages dashboard
# --------------------------------------------------------------------------- #
def build_dashboard(snap, xlsx_rel: str, pdf_rel: str, repo_url: str):
    """Writes docs/index.html (latest) and docs/archive.json (history)."""
    DOCS_DIR.mkdir(exist_ok=True)
    archive_path = DOCS_DIR / "archive.json"
    archive = []
    if archive_path.exists():
        with open(archive_path) as f:
            archive = json.load(f)
    archive = [a for a in archive if a["date"] != snap["report_date"]]
    archive.append({"date": snap["report_date"], "counts": snap["counts"], "xlsx": xlsx_rel, "pdf": pdf_rel})
    archive.sort(key=lambda a: a["date"], reverse=True)
    archive = archive[:10]  # rolling 10-day window
    with open(archive_path, "w") as f:
        json.dump(archive, f, indent=1)

    c = snap["counts"]
    recs = snap["records"]
    blob = repo_url.rstrip("/") + "/blob/main/"

    def row(r):
        sev = r["cvss_severity"]
        return (
            "<tr>"
            f"<td class='pri'>{html.escape(r['priority'])}</td>"
            f"<td><span class='rsn {'new' if r.get('reason')=='New' else 'upd'}'>{html.escape(r.get('reason','New'))}</span></td>"
            f"<td><a href='{r.get('nvd_url') or r.get('advisory_url') or '#'}' target='_blank'>{r['cve_id']}</a></td>"
            f"<td class='num'>{fmt(r['cvss_score'])}</td>"
            f"<td><span class='sev {sev.lower()}'>{sev}</span></td>"
            f"<td>{'<span class=kev>KEV</span>' if r['kev'] else ''}</td>"
            f"<td class='num'>{fmt(r['epss'])}</td>"
            f"<td>{html.escape(r['watchlist_hits'])}</td>"
            f"<td>{html.escape((r['vendor_product'] or r.get('ecosystem_package',''))[:80])}"
            f"{('<br><small>fixed: ' + html.escape(r['patched_version'][:50]) + '</small>') if r.get('patched_version') else ''}</td>"
            f"<td class='src'>{html.escape(', '.join(r.get('sources', [])))}</td>"
            f"<td class='desc'>{html.escape(r['description'][:300])}{'...' if len(r['description']) > 300 else ''}</td>"
            "</tr>"
        )

    history_rows = "".join(
        f"<tr><td>{a['date']}</td><td class='num'>{a['counts']['total']}</td>"
        f"<td class='num'>{a['counts']['kev']}</td><td class='num'>{a['counts']['critical']}</td>"
        f"<td class='num'>{a['counts']['high']}</td>"
        f"<td><a href='{blob}{a['xlsx']}'>Excel</a> &middot; <a href='{blob}{a['pdf']}'>PDF</a></td></tr>"
        for a in archive
    )

    page = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>Daily Vulnerability Report - {snap['report_date']}</title>
<style>
 :root{{--ink:#1a1a1a;--muted:#6b6b6b;--line:#e3e3e3;--navy:#1F3864;--bg:#fafafa}}
 body{{font:14px/1.45 -apple-system,Segoe UI,Roboto,Helvetica,Arial,sans-serif;color:var(--ink);background:var(--bg);margin:0;padding:24px}}
 h1{{font-size:22px;margin:0 0 4px}} .sub{{color:var(--muted);font-size:13px;margin-bottom:18px}}
 .cards{{display:grid;grid-template-columns:repeat(auto-fit,minmax(120px,1fr));gap:10px;margin-bottom:18px}}
 .card{{background:#fff;border:1px solid var(--line);border-radius:8px;padding:12px 14px}}
 .card b{{display:block;font-size:26px;line-height:1.1}} .card span{{color:var(--muted);font-size:12px}}
 .dl a{{display:inline-block;background:var(--navy);color:#fff;text-decoration:none;padding:8px 14px;border-radius:6px;margin-right:8px;font-size:13px}}
 .tools{{margin:18px 0 8px;display:flex;gap:10px;flex-wrap:wrap;align-items:center}}
 input,select{{padding:6px 8px;border:1px solid var(--line);border-radius:6px;font-size:13px}}
 table{{width:100%;border-collapse:collapse;background:#fff;border:1px solid var(--line);font-size:13px}}
 th,td{{padding:7px 8px;border-bottom:1px solid var(--line);vertical-align:top;text-align:left}}
 th{{background:var(--navy);color:#fff;position:sticky;top:0;cursor:pointer;white-space:nowrap}}
 td.num{{text-align:right;font-variant-numeric:tabular-nums}} td.desc{{color:#333;max-width:520px}}
 td.pri{{white-space:nowrap;font-weight:600}}
 .sev{{padding:2px 7px;border-radius:4px;font-size:11px;font-weight:700}}
 .sev.critical{{background:#C00000;color:#fff}} .sev.high{{background:#E97132;color:#fff}}
 .sev.medium{{background:#F2C94C}} .sev.low{{background:#8FBC8F}}
 .kev{{background:#7030A0;color:#fff;padding:2px 7px;border-radius:4px;font-size:11px;font-weight:700}}
 .rsn{{padding:2px 6px;border-radius:4px;font-size:11px;white-space:nowrap}} .rsn.new{{background:#DCE6F5;color:#1F3864}} .rsn.upd{{background:#EEE;color:#555}}
 td.src{{font-size:11px;color:#555;white-space:nowrap}} td small{{color:#2e7d32}}
 details{{margin-top:24px}} summary{{cursor:pointer;font-weight:600}}
 footer{{color:var(--muted);font-size:12px;margin-top:24px}}
</style></head><body>
<h1>Daily Vulnerability Report</h1>
<div class="sub">{snap['report_date']} &nbsp;|&nbsp; window {snap['window_start'][:16]} to {snap['window_end'][:16]} UTC
 &nbsp;|&nbsp; sources: NVD, CISA KEV, EPSS, CISA Vulnrichment, GitHub Advisories, OSV.dev</div>
<div class="cards">
 <div class="card"><b>{c['total']}</b><span>Total</span></div>
 <div class="card"><b style="color:#1F3864">{c.get('new', '')}</b><span>New today</span></div>
 <div class="card"><b style="color:#6b6b6b">{c.get('updated', '')}</b><span>Updated</span></div>
 <div class="card"><b style="color:#7030A0">{c['kev']}</b><span>In CISA KEV</span></div>
 <div class="card"><b style="color:#C00000">{c['critical']}</b><span>Critical</span></div>
 <div class="card"><b style="color:#E97132">{c['high']}</b><span>High</span></div>
 <div class="card"><b>{c['medium']}</b><span>Medium</span></div>
 <div class="card"><b>{c['unscored']}</b><span>Unscored</span></div>
 <div class="card"><b>{c['watchlist']}</b><span>Watchlist</span></div>
 <div class="card"><b>{c.get('ghsa', 0)}</b><span>Packages (GHSA)</span></div>
</div>
<div class="dl"><a href="{blob}{xlsx_rel}">Download Excel</a><a href="{blob}{pdf_rel}">Download PDF</a></div>
<div class="tools">
 <input id="q" placeholder="Filter by CVE, vendor, keyword..." size="36">
 <select id="sev"><option value="">All severities</option><option>CRITICAL</option><option>HIGH</option><option>MEDIUM</option><option>LOW</option></select>
 <select id="rsn"><option value="">New + Updated</option><option value="New">New only</option><option value="Updated">Updated only</option></select>
 <label><input type="checkbox" id="kevonly"> KEV only</label>
 <label><input type="checkbox" id="wlonly"> Watchlist only</label>
 <span id="count" class="sub" style="margin:0"></span>
</div>
<table id="t"><thead><tr>
<th>Priority</th><th>Reason</th><th>ID</th><th>CVSS</th><th>Severity</th><th>KEV</th><th>EPSS</th><th>Watchlist</th><th>Vendor / Product / Package</th><th>Sources</th><th>Description</th>
</tr></thead><tbody>
{''.join(row(r) for r in recs)}
</tbody></table>
<details><summary>Previous reports ({len(archive)})</summary>
<table><thead><tr><th>Date</th><th>Total</th><th>KEV</th><th>Critical</th><th>High</th><th>Files</th></tr></thead>
<tbody>{history_rows}</tbody></table></details>
<footer>Generated {snap['generated_at'][:19].replace('T',' ')} UTC by GitHub Actions. <b>New</b> = published in the window. <b>Updated</b> = published in the last 30 days and received a CVSS score or KEV listing in the window. Older CVEs with metadata-only changes are excluded. Priority: P1 = KEV, EPSS &ge; 0.5 or SSVC active; P2 = CVSS &ge; 9.0, CVSS &ge; 7.0 with EPSS &ge; 0.1, or SSVC poc; P3 = CVSS &ge; 7.0.</footer>
<script>
const q=document.getElementById('q'),sev=document.getElementById('sev'),kev=document.getElementById('kevonly'),wl=document.getElementById('wlonly'),rsn=document.getElementById('rsn');
const rows=[...document.querySelectorAll('#t tbody tr')];
function apply(){{const s=q.value.toLowerCase(),sv=sev.value;let n=0;
 rows.forEach(r=>{{const txt=r.innerText.toLowerCase();const c=r.children;
  const rv=rsn.value;const rt=c[1].innerText.trim();
  const ok=(!s||txt.includes(s))&&(!sv||c[4].innerText===sv)&&(!kev.checked||c[5].innerText==='KEV')&&(!wl.checked||c[7].innerText.trim()!=='')&&(!rv||(rv==='New'?rt==='New':rt!=='New'));
  r.style.display=ok?'':'none';if(ok)n++;}});
 document.getElementById('count').textContent=n+' of '+rows.length+' shown';}}
[q,sev,kev,wl,rsn].forEach(e=>e.addEventListener('input',apply));apply();
document.querySelectorAll('#t th').forEach((th,i)=>th.addEventListener('click',()=>{{
 const asc=th.dataset.asc!=='1';th.dataset.asc=asc?'1':'0';
 const tb=document.querySelector('#t tbody');
 [...tb.rows].sort((a,b)=>{{const x=a.children[i].innerText,y=b.children[i].innerText;const nx=parseFloat(x),ny=parseFloat(y);
  const cmp=(!isNaN(nx)&&!isNaN(ny))?nx-ny:x.localeCompare(y);return asc?cmp:-cmp;}}).forEach(r=>tb.appendChild(r));}}));
</script></body></html>"""
    with open(DOCS_DIR / "index.html", "w") as f:
        f.write(page)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--date")
    ap.add_argument("--repo-url", default=None, help="https://github.com/org/repo (for dashboard links)")
    args = ap.parse_args()

    snap = load_snapshot(args.date)
    d = snap["report_date"]
    y, m = d[:4], d[5:7]
    xlsx = REPORT_DIR / y / m / f"vulns-{d}.xlsx"
    pdf = REPORT_DIR / y / m / f"vulns-{d}.pdf"

    build_xlsx(snap, xlsx)
    build_pdf(snap, pdf)
    repo_url = args.repo_url or f"https://github.com/{__import__('os').getenv('GITHUB_REPOSITORY', 'ORG/REPO')}"
    build_dashboard(snap, str(xlsx.relative_to(ROOT)), str(pdf.relative_to(ROOT)), repo_url)

    # Expose paths to the workflow
    gh_out = __import__("os").getenv("GITHUB_OUTPUT")
    if gh_out:
        with open(gh_out, "a") as f:
            f.write(f"xlsx={xlsx}\npdf={pdf}\nreport_date={d}\n")
            f.write(f"summary=Total {snap['counts']['total']} | KEV {snap['counts']['kev']} | "
                    f"Critical {snap['counts']['critical']} | High {snap['counts']['high']} | "
                    f"Watchlist {snap['counts']['watchlist']}\n")
    print(f"Built {xlsx}\n      {pdf}\n      {DOCS_DIR / 'index.html'}")


if __name__ == "__main__":
    main()
